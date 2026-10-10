"""F11/F13/F15: provenance survives request binding, responses and reports."""
import pytest
from docx import Document
from app.schemas.chat import ChatResponse, GeospatialNDVIResult
from app.agent.report.builder import _render_index, _render_mask
from app.agent.engine.turn_context import TurnToolState, ToolInvocation
from app.agent.types import ToolRunResult


def document_text(doc):
    return "\n".join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])


def test_ndvi_response_schema_preserves_numerical_provenance():
    stats = {"min": .6, "max": .6, "mean": .6, "std": 0., "calibrated": True,
             "total_pixels": 4, "valid_pixels": 2, "nodata_pixels": 2,
             "percentage_denominator": "total_pixels", "math_invalid_pixels": 1,
             "method_version": "ndvi-v2", "resolved_bands": {"red": 4, "nir": 3},
             "calibration": {"method_version": "reflectance-v2", "bands": {"red": {"scale": .0001, "offset": -.1}}}}
    result = GeospatialNDVIResult(type="ndvi", imagery_id="abc123def456", result_url="/results/x.png", stats=stats)
    exported = result.model_dump()["stats"]
    for key, value in stats.items():
        assert exported[key] == value


def test_nonstream_cloud_response_accepts_quality_method_and_null_estimates():
    result = ChatResponse(content="定性结果", model="test", provider="test", geospatial_result={
        "type": "cloud_mask", "imagery_id": "abc123def456", "result_url": "/results/x.png",
        "stats": {"method": "heuristic", "cloud_pct": None, "candidate_bright_pct": 40.},
        "legend": {"label": "疑似亮暗", "classes": [{"code": 4, "name": "未知"}]},
    })
    exported = result.model_dump()["geospatial_result"]
    assert exported["stats"]["cloud_pct"] is None
    assert exported["stats"]["method"] == "heuristic"


def test_report_of_heuristic_has_candidates_and_no_quality_claim():
    doc = Document()
    _render_mask(doc, {"type": "cloud_mask", "stats": {
        "method": "heuristic", "method_validation": "unvalidated", "cloud_pct": None,
        "shadow_pct": None, "clear_pct": None, "candidate_bright_pct": 80., "candidate_dark_pct": 10.,
        "total_pixels": 10, "valid_pixels": 9,
    }}, 1)
    text = document_text(doc)
    assert "疑似高亮占比" in text
    assert "未验证" in text
    assert "不能当作云" in text
    assert "该比例可用于判断" not in text
    assert "晴空占比" not in text


def test_report_names_quality_source_and_decoding_version():
    doc = Document()
    _render_mask(doc, {"type": "cloud_mask", "stats": {
        "method": "quality_mask", "quality_type": "sentinel2_scl", "method_version": "scl-v1",
        "cloud_pct": 20., "shadow_pct": 10., "clear_pct": 40., "snow_pct": 10., "unknown_pct": 20.,
    }}, 1)
    text = document_text(doc)
    assert "sentinel2_scl" in text and "scl-v1" in text
    assert "未知/未分类占比" in text and "雪冰占比" in text
    assert "不代表实测晴空" in text


def test_uncalibrated_index_report_does_not_infer_absolute_vegetation():
    doc = Document()
    _render_index(doc, {"type": "ndvi", "stats": {
        "min": .7, "max": .8, "mean": .75, "std": .05, "calibrated": False,
        "calibration": {"method": "uncalibrated_relative", "method_version": "reflectance-v2"},
        "total_pixels": 4, "valid_pixels": 3,
    }}, 1)
    text = document_text(doc)
    assert "未标定的相对指数" in text
    assert "不能据此给出绝对植被覆盖" in text
    assert "uncalibrated_relative" in text
    assert "3 个有效像素 / 4 个全幅像素" in text


def test_imported_imagery_drops_trusted_bands_from_old_resource():
    state = TurnToolState(trusted_tool_arguments={"calculate_ndvi": {
        "imagery_id": "111111111111", "red_band": 1, "nir_band": 4,
    }})
    state.record(ToolInvocation(name="fetch_scene", arguments={}, result=ToolRunResult(
        tool_context="imported", geospatial_result={"imagery_id": "222222222222"},
        metadata={"band_roles": {"red": 4, "nir": 3}},
    )))
    actual = state.arguments_for("calculate_ndvi", {"imagery_id": "111111111111"})
    assert actual == {"imagery_id": "222222222222"}


@pytest.mark.asyncio
async def test_selected_imagery_binding_does_not_inject_cached_bands(monkeypatch):
    from types import SimpleNamespace
    import app.agent.engine.input as input_module
    from app.schemas.chat import ChatRequest, ChatMessage
    async def provider(*args, **kwargs):
        return SimpleNamespace(messages=[], has_imagery=True)
    async def owns(*args, **kwargs):
        return True
    async def metadata(*args, **kwargs):
        return {"band_roles": {"red": 1, "nir": 4}}
    monkeypatch.setattr(input_module, "build_provider_request_context", provider)
    monkeypatch.setattr(input_module, "user_owns_imagery", owns)
    monkeypatch.setattr(input_module, "get_user_imagery_metadata", metadata)
    request = ChatRequest(messages=[ChatMessage(role="user", content="计算当前影像NDVI")],
                          metadata={"active_imagery_id": "111111111111"})
    turn = await input_module.build_turn_input(request, user_id="owner")
    assert turn.trusted_tool_arguments["calculate_ndvi"] == {"imagery_id": "111111111111"}


@pytest.mark.parametrize("prefix,code", [("invalid_calibration", "invalid_calibration"),
                                         ("artifact_publication_failed", "artifact_publication_failed"),
                                         ("stale_s2_composite", "stale_s2_composite")])
def test_deterministic_compute_failures_have_specific_safe_messages(prefix, code):
    from app.agent.tools.common import classify_rs_tool_failure
    from app.mcp.client import MCPCallError
    actual, message = classify_rs_tool_failure(MCPCallError(prefix + ": backend detail"))
    assert actual == code
    assert "backend detail" not in message
    assert "35%" not in message
    assert "成像于 2022" not in message
