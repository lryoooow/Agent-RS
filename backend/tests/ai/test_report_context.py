from __future__ import annotations

import asyncio
import json
from contextlib import asynccontextmanager
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document

from app.agent.ai_service import AIService, ChatExecutionSetup
from app.agent.engine.service import AutogenTurnOutput
from app.agent.engine.turn_context import ToolInvocation, turn_scope
from app.agent.persistence import PersistenceContext
from app.agent.report.builder import ReportArtifact, ReportError, _select_imagery_analyses
from app.agent.tools.report.runner import run_report
from app.agent.tools.report.schema import ReportArguments
from app.agent.types import AgentTrace, ToolRunResult
from app.auth import conversation_scope, get_current_conversation_id, user_scope
from app.schemas.chat import ChatRequest

IMAGE = 'abcdef012345'
OTHER_IMAGE = '111111111111'


@pytest.fixture
def service(monkeypatch):
    async def prepare(self, request, **kwargs):
        # Persistence may create a chat or replace an unowned/stale client ID.
        cid = (request.metadata or {}).get('verified_id', 'verified-conversation')
        persistence = PersistenceContext(user_id='owner', conversation_id=cid, assistant_message_id='assistant')
        return ChatExecutionSetup(SimpleNamespace(model='fixture', provider='fixture'), persistence, request)
    monkeypatch.setattr(AIService, '_prepare_chat_execution', prepare)
    monkeypatch.setattr(AIService, '_log_response', lambda *a, **kw: None)
    for method in ['save_assistant_response', 'save_streamed_assistant', 'mark_assistant_failed']:
        monkeypatch.setattr('app.agent.ai_service.' + method, AsyncMock(return_value='saved'))
    monkeypatch.setattr('app.agent.ai_service.schedule_after_response', lambda *a, **kw: None)
    return AIService()


@pytest.mark.asyncio
@pytest.mark.parametrize('stream', [False, True])
@pytest.mark.parametrize('client_id', [None, 'stale-unowned-conversation'])
async def test_report_receives_verified_conversation_in_child_task(service, monkeypatch, stream, client_id):
    captured = []
    async def build(**kwargs):
        await asyncio.sleep(0)
        captured.append(kwargs['conversation_id'])
        if not kwargs['conversation_id']:
            raise ReportError('缺少对话上下文', code='no_conversation')
        return ReportArtifact(IMAGE, 'report.docx', f'/api/imagery/{IMAGE}/results/report.docx')
    monkeypatch.setattr('app.agent.tools.report.runner.build_conversation_report', build)
    async def complete(*, user_id, **kwargs):
        with user_scope(user_id):
            result = await asyncio.create_task(run_report(ReportArguments(imagery_id=IMAGE)))
        assert result.error is None
        return AutogenTurnOutput(content=result.tool_context, trace=AgentTrace(enabled=True), geospatial_result=result.geospatial_result)
    async def events(**kwargs):
        result = await complete(**kwargs)
        yield 'delta', result.content
        yield 'final', result
    monkeypatch.setattr('app.agent.ai_service.complete_turn', complete)
    monkeypatch.setattr('app.agent.ai_service.stream_turn_events', events)
    request = ChatRequest(messages=[{'role':'user','content':'生成提取报告'}], conversation_id=client_id)
    with conversation_scope('outer'):
        if stream:
            chunks = [item async for item in service.stream_chat(request)]
            assert any(item.startswith('event: done') for item in chunks), chunks
            done = next(item for item in chunks if item.startswith('event: done'))
            result = json.loads(done.split('data: ',1)[1])['geospatial_result']
        else:
            result = (await service.chat(request)).geospatial_result.model_dump()
        assert get_current_conversation_id() == 'outer'
    assert captured == ['verified-conversation']
    assert result['type'] == 'report' and result['download_url'].endswith('.docx')


@pytest.mark.asyncio
@pytest.mark.parametrize('stream', [False, True])
async def test_conversation_is_restored_after_engine_failure(service, monkeypatch, stream):
    async def fail(**kwargs):
        assert get_current_conversation_id() == 'verified-conversation'
        raise RuntimeError('fixture engine failure')
    async def events(**kwargs):
        await fail(**kwargs)
        yield
    monkeypatch.setattr('app.agent.ai_service.complete_turn', fail)
    monkeypatch.setattr('app.agent.ai_service.stream_turn_events', events)
    request = ChatRequest(messages=[{'role':'user','content':'生成报告'}])
    with conversation_scope('outer'):
        if stream:
            chunks = [item async for item in service.stream_chat(request)]
            assert any(item.startswith('event: error') for item in chunks)
        else:
            with pytest.raises(Exception):
                await service.chat(request)
        assert get_current_conversation_id() == 'outer'


@pytest.mark.asyncio
async def test_stream_close_closes_engine_and_restores_conversation(service, monkeypatch):
    closed = []
    async def events(**kwargs):
        try:
            yield 'delta', '报告处理中'
            await asyncio.sleep(60)
        finally:
            closed.append(get_current_conversation_id())
    monkeypatch.setattr('app.agent.ai_service.stream_turn_events', events)
    with conversation_scope('outer'):
        stream = service.stream_chat(ChatRequest(messages=[{'role':'user','content':'生成报告'}]))
        async for item in stream:
            if item.startswith('event: delta'):
                break
        await stream.aclose()
        assert get_current_conversation_id() == 'outer'
    assert closed == ['verified-conversation']


@pytest.mark.asyncio
async def test_concurrent_reports_do_not_share_conversation_context(service, monkeypatch):
    ready = asyncio.Event()
    started = []
    async def events(*, request, **kwargs):
        cid = request.metadata['verified_id']
        assert get_current_conversation_id() == cid
        started.append(cid)
        if len(started) == 2:
            ready.set()
        await ready.wait()
        await asyncio.sleep(0)
        assert get_current_conversation_id() == cid
        yield 'final', AutogenTurnOutput(content=cid, trace=AgentTrace(enabled=True))
    monkeypatch.setattr('app.agent.ai_service.stream_turn_events', events)
    async def one(cid):
        request = ChatRequest(messages=[{'role':'user','content':'生成报告'}], metadata={'verified_id':cid})
        with conversation_scope('outer-' + cid):
            result = [item async for item in service.stream_chat(request)]
            assert get_current_conversation_id() == 'outer-' + cid
            return result
    results = await asyncio.gather(one('conversation-a'), one('conversation-b'))
    assert all(any(item.startswith('event: done') for item in result) for result in results), results


@pytest.mark.asyncio
async def test_same_turn_report_renders_successful_results_before_persistence(monkeypatch, tmp_path):
    from app.agent.report import builder
    @asynccontextmanager
    async def acquire():
        yield object()
    monkeypatch.setattr(builder, 'fetch_optional_pool', AsyncMock(return_value=SimpleNamespace(acquire=acquire)))
    monkeypatch.setattr(builder, 'get_conversation', AsyncMock(return_value={'id':'conversation'}))
    monkeypatch.setattr(builder, 'list_recent_analysis_results', AsyncMock(return_value=[]))
    monkeypatch.setattr(builder, 'imagery_root', lambda: tmp_path)
    monkeypatch.setattr(builder, 'read_imagery_metadata', lambda _: {'filename':'real-tool-fixture.tif'})
    geo = {
        'type':'instance_segmentation', 'imagery_id':IMAGE, 'model_name':'SAM3',
        'concepts':['building'], 'instance_count':37, 'counts':{'building':37},
        'union_pixels':233837, 'area_m2':338598.87,
    }
    with user_scope('owner'), conversation_scope('conversation'), turn_scope() as state:
        state.record(ToolInvocation('segment_instances',{},ToolRunResult(tool_context='计算完成',geospatial_result=geo)))
        state.record(ToolInvocation('detect_objects',{},ToolRunResult(tool_context='失败',error='inference_failed',geospatial_result={'type':'detection','imagery_id':OTHER_IMAGE})))
        report = await run_report(ReportArguments())
    assert report.error is None
    assert report.geospatial_result['imagery_id'] == IMAGE
    document = Document(tmp_path / IMAGE / 'results' / report.geospatial_result['filename'])
    cells = '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert 'building' in cells and '37' in cells
    paragraphs = '\n'.join(p.text for p in document.paragraphs)
    assert '338,598.87 平方米' in paragraphs and 'SAM3' in paragraphs
    assert OTHER_IMAGE not in '\n'.join(p.text for p in document.paragraphs)


def test_preview_and_report_cards_do_not_select_an_image_or_create_empty_reports():
    analyses = [
        {'geospatial_result':{'type':'segmentation','imagery_id':IMAGE}},
        {'geospatial_result':{'type':'preview','imagery_id':OTHER_IMAGE}},
        {'geospatial_result':{'type':'report','imagery_id':OTHER_IMAGE}},
    ]
    assert _select_imagery_analyses(analyses,None)[0] == IMAGE
    assert _select_imagery_analyses(analyses,OTHER_IMAGE) == (OTHER_IMAGE,[])
    assert _select_imagery_analyses(analyses[1:],None) == (None,[])


def test_mixed_image_payloads_are_selected_independently():
    analyses = [{'geospatial_result':{'type':'segmentation','imagery_id':IMAGE},
                 'tool_result':{'type':'raster_inspect','imagery_id':OTHER_IMAGE}}]
    selected, entries = _select_imagery_analyses(analyses,IMAGE)
    assert selected == IMAGE and entries == [{'geospatial_result':analyses[0]['geospatial_result']}]


@pytest.mark.asyncio
async def test_current_results_cannot_bypass_conversation_ownership(monkeypatch):
    from app.agent.report import builder
    @asynccontextmanager
    async def acquire():
        yield object()
    monkeypatch.setattr(builder,'fetch_optional_pool',AsyncMock(return_value=SimpleNamespace(acquire=acquire)))
    monkeypatch.setattr(builder,'get_conversation',AsyncMock(return_value=None))
    with pytest.raises(ReportError,match='无权访问'):
        await builder.build_conversation_report(conversation_id='foreign',user_id='intruder',current_analyses=[{'geospatial_result':{'type':'segmentation','imagery_id':IMAGE}}])


@pytest.mark.asyncio
async def test_tool_job_keeps_verified_conversation_for_recovery(monkeypatch):
    from app.agent import tool_jobs
    @asynccontextmanager
    async def acquire():
        yield object()
    create = AsyncMock(return_value='job')
    monkeypatch.setattr(tool_jobs,'get_settings',lambda:SimpleNamespace(tool_jobs_enabled=True))
    monkeypatch.setattr(tool_jobs,'fetch_optional_pool',AsyncMock(return_value=SimpleNamespace(acquire=acquire)))
    monkeypatch.setattr(tool_jobs,'create_tool_job',create)
    monkeypatch.setattr(tool_jobs,'mark_job_running',AsyncMock())
    with conversation_scope('verified-conversation'):
        assert await tool_jobs.begin_tool_job(tool_name='generate_report',arguments={},imagery_id=None,user_id='owner') == 'job'
    assert create.call_args.kwargs['conversation_id'] == 'verified-conversation'
