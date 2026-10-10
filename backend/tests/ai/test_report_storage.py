"""F18: reports use the image's actual store, safe images, and durable publication."""
import asyncio
from contextlib import asynccontextmanager
import json
import tempfile
import threading
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document
from fastapi import HTTPException
from PIL import Image

from app.agent.report import builder
from app.agent.report.builder import ReportError, build_conversation_report
from app.api.routes import imagery as imagery_route
from app.auth import user_scope

IMAGE = "abcdef012345"
OTHER = "111111111111"
OWNER = "00000000-0000-4000-8000-000000000011"


def png_bytes(color):
    buffer = BytesIO()
    Image.new("RGB", (96, 64), color).save(buffer, format="PNG")
    return buffer.getvalue()


class MemoryStore:
    def __init__(self):
        self.objects = {}
        self.downloads = []
        self.uploads = []

    async def exists(self, key):
        return key in self.objects

    async def get_to(self, key, destination):
        self.downloads.append(key)
        if key not in self.objects:
            raise FileNotFoundError(key)
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(self.objects[key])

    async def put(self, key, path):
        self.uploads.append(key)
        self.objects[key] = path.read_bytes()

    async def open_stream(self, key):
        if key not in self.objects:
            raise FileNotFoundError(key)
        async def chunks():
            yield self.objects[key]
        return chunks()


@pytest.fixture
def storage_case(monkeypatch, tmp_path):
    root = tmp_path / "imagery"
    private = tmp_path / "private"
    private.mkdir()
    def temporary_directory(**kwargs):
        kwargs.setdefault("dir", private)
        return tempfile.TemporaryDirectory(**kwargs)
    monkeypatch.setattr(builder, "tempfile", SimpleNamespace(TemporaryDirectory=temporary_directory))
    meta = {"filename": "owned-image.tif", "owner_user_id": OWNER,
            "preview_url": f"/api/imagery/{IMAGE}/results/preview.png"}
    row = {"imagery_id": IMAGE, "owner_user_id": OWNER,
           "storage_backend": "minio", "metadata": meta}
    analyses = [{"geospatial_result": {"type": "ndvi", "imagery_id": IMAGE,
                  "result_url": f"/api/imagery/{IMAGE}/results/ndvi_owned.png",
                  "stats": {"min": .2, "max": .8, "mean": .5, "std": .1}}}]
    @asynccontextmanager
    async def acquire():
        yield object()
    monkeypatch.setattr(builder, "fetch_optional_pool", AsyncMock(return_value=SimpleNamespace(acquire=acquire)))
    monkeypatch.setattr(builder, "get_conversation", AsyncMock(return_value={"id": "owned-conversation"}))
    monkeypatch.setattr(builder, "list_recent_analysis_results", AsyncMock(return_value=analyses))
    monkeypatch.setattr(builder, "db_get_imagery", AsyncMock(return_value=row), raising=False)
    monkeypatch.setattr(builder, "imagery_root", lambda: root)
    monkeypatch.setattr(builder, "read_imagery_metadata", lambda _: None)
    store = MemoryStore()
    store.objects[f"{IMAGE}/results/preview.png"] = png_bytes("green")
    store.objects[f"{IMAGE}/results/ndvi_owned.png"] = png_bytes("blue")
    monkeypatch.setattr(builder, "object_store_for", lambda backend: store, raising=False)
    monkeypatch.setattr(imagery_route, "_imagery_root", lambda: root)
    monkeypatch.setattr(imagery_route, "_effective_backend", AsyncMock(return_value="minio"))
    monkeypatch.setattr(imagery_route, "_ensure_owned_for_mutation", AsyncMock())
    monkeypatch.setattr(imagery_route, "object_store_for", lambda backend: store)
    return SimpleNamespace(root=root, private=private, meta=meta, row=row, analyses=analyses, store=store)


@pytest.mark.asyncio
async def test_minio_report_is_downloadable_and_embeds_owned_figures(storage_case):
    report = await build_conversation_report(conversation_id="owned-conversation", user_id=OWNER)
    with user_scope(OWNER):
        try:
            response = await imagery_route.get_result_file(IMAGE, report.filename)
        except HTTPException as exc:
            pytest.fail(f"Report returned a download URL, but its real object-store download is HTTP {exc.status_code}")
    content = b"".join([part async for part in response.body_iterator])
    document = Document(BytesIO(content))
    assert len(document.inline_shapes) == 2
    assert "owned-image.tif" in " ".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert response.media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert storage_case.store.uploads == [f"{IMAGE}/results/{report.filename}"]
    assert all(key.endswith(".png") for key in storage_case.store.downloads)
    assert not list(storage_case.root.rglob("*.docx"))


@pytest.mark.asyncio
async def test_explicit_imagery_id_uses_same_minio_store(storage_case):
    report = await build_conversation_report(conversation_id="owned-conversation", user_id=OWNER, imagery_id=IMAGE)
    assert f"{IMAGE}/results/{report.filename}" in storage_case.store.objects


@pytest.mark.asyncio
async def test_report_tool_without_imagery_id_uses_owned_result_store(storage_case):
    from app.agent.tools.report.runner import run_report
    from app.agent.tools.report.schema import ReportArguments
    from app.auth import conversation_scope
    with user_scope(OWNER), conversation_scope("owned-conversation"):
        result = await run_report(ReportArguments())
    assert result.error is None
    assert result.geospatial_result["imagery_id"] == IMAGE
    assert f"{IMAGE}/results/{result.geospatial_result['filename']}" in storage_case.store.objects


@pytest.mark.asyncio
async def test_report_http_without_imagery_id_publishes_before_saving_card(storage_case, monkeypatch):
    from app.api.routes import report as route
    saved = AsyncMock()
    monkeypatch.setattr(route, "save_assistant_response", saved)
    with user_scope(OWNER):
        report = await route.create_report(route.ReportCreateRequest(conversation_id="owned-conversation"))
    assert saved.await_count == 1
    assert f"{IMAGE}/results/{report.filename}" in storage_case.store.objects


@pytest.mark.asyncio
async def test_report_publish_failure_returns_no_download_card(storage_case, monkeypatch):
    from app.api.routes import report as route
    saved = AsyncMock()
    monkeypatch.setattr(route, "save_assistant_response", saved)
    async def failed_put(key, path):
        raise OSError("synthetic private publication detail")
    monkeypatch.setattr(storage_case.store, "put", failed_put)
    with user_scope(OWNER), pytest.raises(HTTPException) as caught:
        await route.create_report(route.ReportCreateRequest(conversation_id="owned-conversation"))
    assert caught.value.status_code == 502
    assert caught.value.detail["code"] == "report_publish_failed"
    assert "synthetic private publication detail" not in caught.value.detail["message"]
    saved.assert_not_awaited()
    assert not list(storage_case.private.iterdir())
    assert not any(key.endswith(".docx") for key in storage_case.store.objects)


@pytest.mark.asyncio
async def test_image_owner_is_checked_even_when_conversation_is_owned(storage_case):
    storage_case.row["owner_user_id"] = "foreign-user"
    with pytest.raises(ReportError) as caught:
        await build_conversation_report(conversation_id="owned-conversation", user_id=OWNER)
    assert caught.value.code == "imagery_forbidden"
    assert storage_case.store.downloads == storage_case.store.uploads == []
    assert not list(storage_case.private.iterdir())


@pytest.mark.asyncio
async def test_conversation_owner_denial_happens_before_image_lookup(storage_case, monkeypatch):
    monkeypatch.setattr(builder, "get_conversation", AsyncMock(return_value=None))
    with pytest.raises(ReportError) as caught:
        await build_conversation_report(conversation_id="foreign-conversation", user_id=OWNER)
    assert caught.value.code == "conversation_forbidden"
    builder.db_get_imagery.assert_not_awaited()
    assert storage_case.store.downloads == storage_case.store.uploads == []


@pytest.mark.asyncio
async def test_minio_image_row_is_used_when_global_storage_switches_local(storage_case, monkeypatch):
    monkeypatch.setenv("STORAGE_BACKEND", "local")
    report = await build_conversation_report(conversation_id="owned-conversation", user_id=OWNER)
    assert f"{IMAGE}/results/{report.filename}" in storage_case.store.objects
    assert not list(storage_case.root.rglob("*.docx"))


@pytest.mark.asyncio
async def test_local_image_row_is_used_when_global_storage_switches_minio(storage_case, monkeypatch):
    monkeypatch.setenv("STORAGE_BACKEND", "minio")
    storage_case.row["storage_backend"] = "local"
    results = storage_case.root / IMAGE / "results"
    results.mkdir(parents=True)
    (results.parent / "metadata.json").write_text(json.dumps(storage_case.meta), encoding="utf-8")
    for name in ("preview.png", "ndvi_owned.png"):
        (results / name).write_bytes(storage_case.store.objects[f"{IMAGE}/results/{name}"])
    report = await build_conversation_report(conversation_id="owned-conversation", user_id=OWNER)
    document = Document(results / report.filename)
    assert len(document.inline_shapes) == 2
    assert storage_case.store.downloads == storage_case.store.uploads == []
    assert [path.name for path in results.iterdir() if path.is_dir()] == []
    monkeypatch.setattr(imagery_route, "_effective_backend", AsyncMock(return_value="local"))
    with user_scope(OWNER):
        response = await imagery_route.get_result_file(IMAGE, report.filename)
    assert Path(response.path) == results / report.filename


@pytest.mark.asyncio
async def test_missing_optional_images_are_not_fabricated(storage_case):
    storage_case.store.objects.clear()
    report = await build_conversation_report(conversation_id="owned-conversation", user_id=OWNER)
    document = Document(BytesIO(storage_case.store.objects[f"{IMAGE}/results/{report.filename}"]))
    assert len(document.inline_shapes) == 0
    assert storage_case.store.downloads == []


@pytest.mark.asyncio
async def test_foreign_image_url_is_never_downloaded(storage_case):
    storage_case.meta["preview_url"] = "https://untrusted.invalid/private.png"
    storage_case.analyses[0]["geospatial_result"]["result_url"] = f"/api/imagery/{OTHER}/results/private.png"
    storage_case.store.objects[f"{OTHER}/results/private.png"] = png_bytes("red")
    report = await build_conversation_report(conversation_id="owned-conversation", user_id=OWNER)
    document = Document(BytesIO(storage_case.store.objects[f"{IMAGE}/results/{report.filename}"]))
    assert len(document.inline_shapes) == 1  # Only this owner's default preview.
    assert storage_case.store.downloads == [f"{IMAGE}/results/preview.png"]


@pytest.mark.parametrize("value", [
    "C:/private/a.png", "file:///private/a.png", "/private/a.png", "../private.png",
    "https://untrusted.invalid/private.png", "http://[",
    f"/api/imagery/{OTHER}/results/a.png",
    f"/api/imagery/{IMAGE}/results/%2e%2e%2fprivate.png",
    f"/api/imagery/{IMAGE}/results/%2e%2e%5cprivate.png",
    f"/api/imagery/{IMAGE}/results/a.svg",
    f"/api/imagery/{IMAGE}/results/%00.png",
])
def test_artifact_names_cannot_select_other_paths_or_external_images(value):
    assert builder._artifact_name(IMAGE, value, image_only=True) is None


@pytest.mark.asyncio
async def test_failed_render_never_leaves_a_public_partial_report(storage_case, monkeypatch):
    def failed_render(output, **kwargs):
        output.write_bytes(b"partial ZIP")
        raise OSError("synthetic render error")
    monkeypatch.setattr(builder, "_render_report_docx", failed_render)
    with pytest.raises(ReportError) as caught:
        await build_conversation_report(conversation_id="owned-conversation", user_id=OWNER)
    assert caught.value.code == "render_failed"
    assert not list(storage_case.private.iterdir())
    assert not list(storage_case.root.rglob("*.docx"))
    assert storage_case.store.uploads == []


@pytest.mark.asyncio
async def test_local_report_remains_private_until_renderer_finishes(storage_case, monkeypatch):
    storage_case.row["storage_backend"] = "local"
    started = asyncio.Event()
    release = threading.Event()
    loop = asyncio.get_running_loop()
    original = builder._render_report_docx
    captured = {}
    def blocked_render(output, **kwargs):
        captured["output"] = output
        loop.call_soon_threadsafe(started.set)
        assert release.wait(5)
        original(output, **kwargs)
    monkeypatch.setattr(builder, "_render_report_docx", blocked_render)
    task = asyncio.create_task(build_conversation_report(conversation_id="owned-conversation", user_id=OWNER))
    try:
        await asyncio.wait_for(started.wait(), 5)
        results = storage_case.root / IMAGE / "results"
        assert not list(results.glob("*.docx"))
        assert captured["output"].parent.name.startswith(".report_")
    finally:
        release.set()
    report = await task
    assert (results / report.filename).is_file()
    assert not captured["output"].exists()


@pytest.mark.asyncio
async def test_cancelled_render_keeps_workspace_until_worker_releases_it(storage_case, monkeypatch):
    started = asyncio.Event()
    finished = asyncio.Event()
    release = threading.Event()
    loop = asyncio.get_running_loop()
    captured = {}
    def blocked_render(output, **kwargs):
        captured["output"] = output
        loop.call_soon_threadsafe(started.set)
        assert release.wait(5)
        output.write_bytes(b"completed private render")
        loop.call_soon_threadsafe(finished.set)
    monkeypatch.setattr(builder, "_render_report_docx", blocked_render)
    task = asyncio.create_task(build_conversation_report(conversation_id="owned-conversation", user_id=OWNER))
    try:
        await asyncio.wait_for(started.wait(), 5)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert captured["output"].parent.exists()
        assert storage_case.store.uploads == []
    finally:
        release.set()
    await asyncio.wait_for(finished.wait(), 5)
    for _ in range(10):
        await asyncio.sleep(0)
        if not captured["output"].parent.exists():
            break
    assert not captured["output"].parent.exists()
    assert storage_case.store.uploads == []


@pytest.mark.asyncio
async def test_cancelled_publish_keeps_complete_file_until_transfer_finishes(storage_case, monkeypatch):
    started = asyncio.Event()
    finished = asyncio.Event()
    release = asyncio.Event()
    original = storage_case.store.put
    captured = {}
    async def blocked_put(key, path):
        captured["output"] = path
        started.set()
        await release.wait()
        await original(key, path)
        finished.set()
    monkeypatch.setattr(storage_case.store, "put", blocked_put)
    task = asyncio.create_task(build_conversation_report(conversation_id="owned-conversation", user_id=OWNER))
    await asyncio.wait_for(started.wait(), 5)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert captured["output"].exists()
    release.set()
    await asyncio.wait_for(finished.wait(), 5)
    await asyncio.sleep(0)
    assert not captured["output"].parent.exists()


@pytest.mark.asyncio
async def test_same_image_parallel_reports_do_not_reuse_files(storage_case):
    reports = await asyncio.gather(*[
        build_conversation_report(conversation_id="owned-conversation", user_id=OWNER) for _ in range(2)
    ])
    assert reports[0].filename != reports[1].filename
    assert len(storage_case.store.uploads) == 2
    for report in reports:
        document = Document(BytesIO(storage_case.store.objects[f"{IMAGE}/results/{report.filename}"]))
        assert len(document.inline_shapes) == 2
    assert not list(storage_case.private.iterdir())


@pytest.mark.asyncio
async def test_backend_database_failure_does_not_fall_back_to_local_cache(storage_case, monkeypatch):
    monkeypatch.setattr(builder, "db_get_imagery", AsyncMock(side_effect=OSError("synthetic database failure")))
    with pytest.raises(ReportError) as caught:
        await build_conversation_report(conversation_id="owned-conversation", user_id=OWNER)
    assert caught.value.code == "storage_unavailable"
    assert storage_case.store.downloads == storage_case.store.uploads == []


async def seed_pg_report_image(pool, *, backend="minio", image_owner=None):
    from app.core.settings import Settings
    from app.db.repositories._pg import conversation, identity, imagery, message
    settings = Settings(_env_file=None)
    owner = settings.default_user_id
    meta = {"filename": "pg-owned.tif", "owner_user_id": image_owner or owner,
            "preview_url": f"/api/imagery/{IMAGE}/results/preview.png"}
    async with pool.acquire() as conn:
        await identity.ensure_default_identity(conn, settings)
        cid = await conversation.create_conversation(conn, user_id=owner, settings=settings, title="Storage report test")
        await message.append_message(conn, conversation_id=cid, role="assistant", content="NDVI computed",
            metadata={"finish_reason": "stop", "geospatial_result": {
                "type": "ndvi", "imagery_id": IMAGE,
                "result_url": f"/api/imagery/{IMAGE}/results/ndvi_owned.png",
                "stats": {"mean": .5, "min": .2, "max": .8, "std": .1}}})
        await imagery.insert_imagery(conn, imagery_id=IMAGE, owner_user_id=image_owner or owner,
            filename=meta["filename"], storage_backend=backend, metadata=meta)
    return cid, owner


@pytest.mark.asyncio(loop_scope="session")
@pytest.mark.parametrize("entrypoint", ["tool", "http"])
async def test_real_pg_minio_report_download_and_owner_guard(pg_pool, monkeypatch, tmp_path, entrypoint):
    from app.api.routes import report as report_route
    from app.agent.tools.report.runner import run_report
    from app.agent.tools.report.schema import ReportArguments
    from app.auth import conversation_scope
    from app.db.repositories.message import list_conversation_messages
    cid, owner = await seed_pg_report_image(pg_pool)
    store = MemoryStore()
    store.objects[f"{IMAGE}/results/preview.png"] = png_bytes("green")
    store.objects[f"{IMAGE}/results/ndvi_owned.png"] = png_bytes("blue")
    pool_fn = AsyncMock(return_value=pg_pool)
    monkeypatch.setattr(builder, "fetch_optional_pool", pool_fn)
    monkeypatch.setattr(imagery_route, "fetch_optional_pool", pool_fn)
    monkeypatch.setattr("app.agent.persistence.fetch_optional_pool", pool_fn)
    monkeypatch.setattr(builder, "object_store_for", lambda backend: store)
    monkeypatch.setattr(imagery_route, "object_store_for", lambda backend: store)
    monkeypatch.setattr(builder, "imagery_root", lambda: tmp_path / "imagery")
    monkeypatch.setattr(imagery_route, "_imagery_root", lambda: tmp_path / "imagery")
    monkeypatch.setenv("STORAGE_BACKEND", "local")
    with user_scope(owner), conversation_scope(cid):
        if entrypoint == "http":
            report = await report_route.create_report(report_route.ReportCreateRequest(conversation_id=cid))
            filename = report.filename
            async with pg_pool.acquire() as conn:
                saved = await list_conversation_messages(conn, conversation_id=cid, user_id=owner)
            assert saved[-1]["metadata_json"]["geospatial_result"]["download_url"] == report.download_url
        else:
            result = await run_report(ReportArguments())
            assert result.error is None
            filename = result.geospatial_result["filename"]
        response = await imagery_route.get_result_file(IMAGE, filename)
    content = b"".join([chunk async for chunk in response.body_iterator])
    document = Document(BytesIO(content))
    assert len(document.inline_shapes) == 2
    assert store.downloads == [f"{IMAGE}/results/ndvi_owned.png", f"{IMAGE}/results/preview.png"]
    assert not list((tmp_path / "imagery").rglob("*.docx"))
    with user_scope("00000000-0000-4000-8000-0000000000ff"), pytest.raises(HTTPException) as caught:
        await imagery_route.get_result_file(IMAGE, filename)
    assert caught.value.status_code == 404


@pytest.mark.asyncio(loop_scope="session")
async def test_real_pg_owned_conversation_cannot_render_foreign_image(pg_pool, monkeypatch, tmp_path):
    cid, owner = await seed_pg_report_image(pg_pool, image_owner="00000000-0000-4000-8000-0000000000ff")
    store = MemoryStore()
    monkeypatch.setattr(builder, "fetch_optional_pool", AsyncMock(return_value=pg_pool))
    monkeypatch.setattr(builder, "object_store_for", lambda backend: store)
    monkeypatch.setattr(builder, "imagery_root", lambda: tmp_path / "imagery")
    with pytest.raises(ReportError) as caught:
        await build_conversation_report(conversation_id=cid, user_id=owner)
    assert caught.value.code == "imagery_forbidden"
    assert store.downloads == store.uploads == []
    assert not list((tmp_path / "imagery").rglob("*.docx"))


@pytest.mark.asyncio
@pytest.mark.parametrize("backend", ["local", "minio"])
async def test_existing_report_filename_is_never_overwritten(storage_case, monkeypatch, backend):
    from datetime import datetime, timezone
    fixed = datetime(2026, 10, 9, tzinfo=timezone.utc)
    monkeypatch.setattr(builder, "datetime", SimpleNamespace(now=lambda _: fixed))
    monkeypatch.setattr(builder, "uuid", SimpleNamespace(uuid4=lambda: SimpleNamespace(hex="f" * 32)))
    storage_case.row["storage_backend"] = backend
    filename = "report_20261009_000000_ffffffffffffffff.docx"
    if backend == "local":
        results = storage_case.root / IMAGE / "results"
        results.mkdir(parents=True)
        (results / filename).write_bytes(b"existing immutable report")
    else:
        storage_case.store.objects[f"{IMAGE}/results/{filename}"] = b"existing immutable report"
    with pytest.raises(ReportError) as caught:
        await build_conversation_report(conversation_id="owned-conversation", user_id=OWNER)
    assert caught.value.code == "report_publish_failed"
    if backend == "local":
        assert (results / filename).read_bytes() == b"existing immutable report"
        assert not any(path.is_dir() for path in results.iterdir())
    else:
        assert storage_case.store.objects[f"{IMAGE}/results/{filename}"] == b"existing immutable report"
        assert storage_case.store.uploads == []
        assert not list(storage_case.private.iterdir())


@pytest.mark.asyncio
async def test_legacy_image_uses_owned_local_metadata(storage_case, monkeypatch):
    monkeypatch.setattr(builder, "db_get_imagery", AsyncMock(return_value=None))
    monkeypatch.setattr(builder, "read_imagery_metadata", lambda _: storage_case.meta)
    report = await build_conversation_report(conversation_id="owned-conversation", user_id=OWNER)
    assert (storage_case.root / IMAGE / "results" / report.filename).is_file()
    assert storage_case.store.uploads == storage_case.store.downloads == []


@pytest.mark.asyncio
async def test_legacy_image_with_foreign_owner_is_not_read(storage_case, monkeypatch):
    monkeypatch.setattr(builder, "db_get_imagery", AsyncMock(return_value=None))
    monkeypatch.setattr(builder, "read_imagery_metadata", lambda _: {"owner_user_id": "foreign"})
    with pytest.raises(ReportError) as caught:
        await build_conversation_report(conversation_id="owned-conversation", user_id=OWNER)
    assert caught.value.code == "imagery_forbidden"
    assert storage_case.store.uploads == storage_case.store.downloads == []
