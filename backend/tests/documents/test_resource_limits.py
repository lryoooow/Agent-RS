from __future__ import annotations
import asyncio
import json
import struct
import subprocess
import sys
import zipfile
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import pytest
from app.core.settings import Settings
from app.documents.parser import DocumentParseError, parse_uploaded_document
from app.documents.isolated import parse_document_file


def archive(payload):
    target=BytesIO()
    with zipfile.ZipFile(target,"w",zipfile.ZIP_DEFLATED) as out: out.writestr("word/document.xml",payload)
    return bytearray(target.getvalue())


def test_forged_expansion_size_is_checked_against_actual_stream():
    data=archive(b"a"*200_000)
    central=data.index(b"PK\x01\x02")
    struct.pack_into("<I",data,central+24,10)  # underdeclare expanded size
    with pytest.raises(DocumentParseError) as error:
        parse_uploaded_document(filename="fake.docx",content_type=None,data=bytes(data),title=None,
            settings=Settings(_env_file=None,document_office_max_total_uncompressed_bytes=50_000,
                document_office_max_entry_compression_ratio=10_000))
    assert error.value.code=="OFFICE_ARCHIVE_TOO_LARGE"


def test_forged_entry_count_rejected_before_zipinfo_allocation(monkeypatch):
    data=archive(b"ok"); end=data.rfind(b"PK\x05\x06")
    struct.pack_into("<HH",data,end+8,0,0)
    monkeypatch.setattr(zipfile,"ZipFile",lambda *a,**k: pytest.fail("must reject before allocation"))
    with pytest.raises(DocumentParseError,match="Could not parse"):
        parse_uploaded_document(filename="fake.docx",content_type=None,data=bytes(data),title=None,settings=Settings(_env_file=None))


def test_corrupt_crc_rejected_before_format_parser():
    data=archive(b"ok"); central=data.index(b"PK\x01\x02")
    struct.pack_into("<I",data,central+16,0)
    with pytest.raises(DocumentParseError) as error:
        parse_uploaded_document(filename="bad.docx",content_type=None,data=bytes(data),title=None,settings=Settings(_env_file=None))
    assert error.value.code=="DOCX_PARSE_FAILED"


@pytest.mark.parametrize("rows,cols",[(4,1),(1,4)])
def test_sheet_dimensions_do_not_drive_unbounded_iteration(rows,cols):
    from openpyxl import Workbook
    workbook=Workbook(); workbook.active.cell(rows,cols,"x")
    target=BytesIO(); workbook.save(target)
    with pytest.raises(DocumentParseError) as error:
        parse_uploaded_document(filename="large.xlsx",content_type=None,data=target.getvalue(),title=None,
            settings=Settings(_env_file=None,document_max_sheet_rows=3,document_max_sheet_columns=3))
    assert error.value.code=="WORKSHEET_TOO_LARGE"


def test_text_limit_boundary_and_overflow():
    settings=Settings(_env_file=None,document_max_text_chars=3)
    assert parse_uploaded_document(filename="a.txt",content_type=None,data="中文A".encode(),title=None,settings=settings).content=="中文A"
    with pytest.raises(DocumentParseError) as error:
        parse_uploaded_document(filename="a.txt",content_type=None,data="中文AB".encode(),title=None,settings=settings)
    assert error.value.code=="DOCUMENT_TEXT_TOO_LARGE"


@pytest.mark.asyncio
async def test_real_isolated_parser_preserves_structure_and_cleans_workspace(tmp_path):
    from docx import Document
    doc=Document(); doc.add_heading("章节",1); doc.add_paragraph("遥感正文")
    path=tmp_path/"source.docx"; doc.save(path)
    result=await parse_document_file(temp_path=path,filename="source.docx",content_type=None,title=None,settings=Settings(_env_file=None))
    assert "# 章节" in result.content and "遥感正文" in result.content
    assert path.exists()  # input ownership remains with ingestion caller


@pytest.mark.asyncio
async def test_real_parser_timeout_returns_bounded_error(tmp_path):
    path=tmp_path/"source.txt"; path.write_text("small",encoding="utf-8")
    with pytest.raises(DocumentParseError) as error:
        await parse_document_file(temp_path=path,filename="source.txt",content_type=None,title=None,
            settings=Settings(_env_file=None,document_parse_timeout_seconds=.001))
    assert error.value.code=="DOCUMENT_PARSE_TIMEOUT" and path.exists()


@pytest.mark.asyncio
async def test_cancellation_reaps_worker_before_return(monkeypatch,tmp_path):
    started,killed,reaped=asyncio.Event(),asyncio.Event(),asyncio.Event()
    class Process:
        returncode=None
        async def wait(self):
            started.set(); await killed.wait(); self.returncode=-9; reaped.set(); return -9
        def kill(self): killed.set()
    async def spawn(*args,**kwargs): return Process()
    monkeypatch.setattr(asyncio,"create_subprocess_exec",spawn)
    task=asyncio.create_task(parse_document_file(temp_path=tmp_path/"source",filename="source.docx",
        content_type=None,title=None,settings=Settings(_env_file=None)))
    await started.wait(); task.cancel()
    with pytest.raises(asyncio.CancelledError): await task
    assert killed.is_set() and reaped.is_set()


def test_actual_os_memory_limit_rejects_oversized_allocation():
    script="from app.documents.worker import limit_memory; limit_memory(64*1024*1024);\ntry: data=bytearray(128*1024*1024)\nexcept MemoryError: print('limited')\nelse: print('unlimited')"
    import os
    env=os.environ.copy(); env["PYTHONPATH"]=str(Path(__file__).resolve().parents[2])
    result=subprocess.run([sys.executable,"-B","-c",script],env=env,capture_output=True,text=True,timeout=20)
    assert result.returncode==0 and result.stdout.strip()=="limited"


@pytest.mark.asyncio
async def test_parser_concurrency_is_bounded_without_blocking_health_task(monkeypatch,tmp_path):
    import app.documents.isolated as module
    started=0; active=0; peak=0
    release=asyncio.Event()
    class Process:
        returncode=None
        def __init__(self,output): self.output=output
        async def wait(self):
            nonlocal active
            await release.wait()
            if self.returncode is None:
                self.output.write_text(json.dumps({"document":{"title":"t","content":"text","doc_type":"txt","metadata":{}}}),encoding="utf-8")
                self.returncode=0; active-=1
            return self.returncode
        def kill(self): self.returncode=-9; release.set()
    async def spawn(*args,**kwargs):
        nonlocal started,active,peak
        started+=1; active+=1; peak=max(peak,active)
        return Process(Path(args[-1]))
    monkeypatch.setattr(asyncio,"create_subprocess_exec",spawn)
    settings=Settings(_env_file=None,document_parse_max_concurrent=2)
    tasks=[asyncio.create_task(parse_document_file(temp_path=tmp_path/"source",filename="s.txt",content_type=None,title=None,settings=settings)) for i in range(3)]
    for _ in range(100):
        if started==2:break
        await asyncio.sleep(.001)
    assert started==2 and active==2
    async def health(): await asyncio.sleep(0); return "alive"
    assert await asyncio.wait_for(health(),.1)=="alive"
    release.set()
    results=await asyncio.wait_for(asyncio.gather(*tasks),2)
    assert peak==2 and started==3 and [r.content for r in results]==["text"]*3


@pytest.mark.asyncio
async def test_cancelled_ingestion_persists_terminal_state_and_removes_input(monkeypatch,tmp_path):
    import app.api.routes.documents as module
    from unittest.mock import AsyncMock
    class Pool:
        def acquire(self):return self
        async def __aenter__(self):return object()
        async def __aexit__(self,*args):return False
    started=asyncio.Event()
    async def parse(**kwargs): started.set(); await asyncio.Event().wait()
    monkeypatch.setattr(module,"_require_document_db",AsyncMock(return_value=Pool()))
    monkeypatch.setattr(module,"update_ingest_job",AsyncMock())
    monkeypatch.setattr(module,"parse_document_file",parse)
    terminal=AsyncMock();monkeypatch.setattr(module,"_fail_ingest_job",terminal)
    path=tmp_path/"source.txt";path.write_text("private source")
    task=asyncio.create_task(module._run_document_ingest_job(job_id="job",temp_path=path,filename="s.txt",content_type=None,title=None,metadata={},user_id="owner"))
    await started.wait();task.cancel()
    with pytest.raises(asyncio.CancelledError):await asyncio.wait_for(task,2)
    terminal.assert_awaited_once_with("job","DOCUMENT_INGEST_CANCELLED","Document ingestion was cancelled.")
    assert not path.exists()


@pytest.mark.asyncio
async def test_isolated_excel_limits_native_threads_in_worker_only(tmp_path, monkeypatch):
    import os
    from openpyxl import Workbook
    from app.core.settings import Settings
    from app.documents.isolated import parse_document_file
    from app.documents import isolated
    workbook = Workbook()
    workbook.active.append(["worker thread verification", 123])
    path = tmp_path / "worker.xlsx"
    workbook.save(path)
    monkeypatch.setenv("OPENBLAS_NUM_THREADS", "64")
    launch = isolated.asyncio.create_subprocess_exec
    child_environments = []
    async def captured(*args, **kwargs):
        child_environments.append(kwargs["env"])
        return await launch(*args, **kwargs)
    monkeypatch.setattr(isolated.asyncio, "create_subprocess_exec", captured)
    result = await parse_document_file(temp_path=path, filename=path.name,
        content_type=None, title="check", settings=Settings(_env_file=None))
    assert "worker thread verification" in result.content
    assert os.environ["OPENBLAS_NUM_THREADS"] == "64"
    assert all(child_environments[0][key] == "1" for key in
        ("OPENBLAS_NUM_THREADS", "OMP_NUM_THREADS", "MKL_NUM_THREADS", "NUMEXPR_NUM_THREADS"))
